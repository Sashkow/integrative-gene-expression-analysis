#!/usr/bin/env python3
"""
Build Excel DEG table for yehor conference with 1_2, 1_2_m, 1_2_f, and Prater columns.

- Italic: near-miss genes
- Bold: candidate sex-specific (in one sex only AND not near-miss in opposite)

Usage:
    python3 scripts/yehor_conference_2026/create_deg_table.py
"""

import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "articles/yehor_conference_2026/data/phase2b_1_2_yehor_6ds_no_37653_no_22490_enriched_sashko"
SEX_DIR = f"{BASE}/sex_stratified"
VENN_DIR = f"{SEX_DIR}/venn_m_vs_f"
PRATER_PATH = "articles/yehor_conference_2026/references/Prater_2021_RNA-Seq_first-second_trimester_transition_supp_tables.xlsx"
OUTPUT = "articles/yehor_conference_2026/data/deg_table_sex_stratified.xlsx"

FDR_THRESH = 0.05
LOGFC_THRESH = 1.0
NEAR_MISS_FDR = 0.10
NEAR_MISS_LOGFC = 0.8


def load_de_table(path):
    with open(path) as f:
        reader = csv.DictReader(f, delimiter="\t")
        return {r["gene"]: r for r in reader}


def load_sig_genes(path):
    with open(path) as f:
        reader = csv.DictReader(f, delimiter="\t")
        return set(r["gene"] for r in reader)


def load_prater():
    wb = openpyxl.load_workbook(PRATER_PATH, read_only=True)
    ws = wb["T1 DEGs_results_table_l2fc1"]
    rows = list(ws.iter_rows(values_only=True))
    result = {}
    for row in rows[1:]:
        ensembl, baseMean, l2fc, lfcSE, stat, pval, padj, gene_name, entrez, desc = row
        if entrez is not None and str(entrez) != "NA":
            try:
                result[str(int(entrez))] = {
                    "symbol": gene_name,
                    "log2FC": float(l2fc),
                    "padj": float(padj) if padj and str(padj) != "NA" else None,
                }
            except (ValueError, TypeError):
                pass
    wb.close()
    return result


def load_annotations():
    path = f"{VENN_DIR}/all_groups_annotated.tsv"
    with open(path) as f:
        reader = csv.DictReader(f, delimiter="\t")
        return {r["ENTREZID"]: r for r in reader}


def load_symbols():
    import subprocess
    result = {}
    try:
        ann = load_annotations()
        for eid, r in ann.items():
            result[eid] = {
                "symbol": r.get("SYMBOL", ""),
                "genename": r.get("GENENAME", ""),
                "chr": r.get("chromosome", ""),
            }
    except Exception:
        pass
    return result


def is_sig(de_row):
    return (float(de_row["adj.P.Val"]) < FDR_THRESH and
            abs(float(de_row["logFC"])) >= LOGFC_THRESH)


def is_near_miss(de_row):
    return (float(de_row["adj.P.Val"]) < NEAR_MISS_FDR and
            abs(float(de_row["logFC"])) >= NEAR_MISS_LOGFC)


def main():
    de_12 = load_de_table(f"{BASE}/difexp_softimpute_combat_ref.tsv")
    de_m = load_de_table(f"{SEX_DIR}/difexp_males_1t_vs_2t.tsv")
    de_f = load_de_table(f"{SEX_DIR}/difexp_females_1t_vs_2t.tsv")

    sig_12 = load_sig_genes(f"{BASE}/difexp_significant_softimpute_combat_ref.tsv")
    sig_m = load_sig_genes(f"{SEX_DIR}/difexp_significant_males_1t_vs_2t.tsv")
    sig_f = load_sig_genes(f"{SEX_DIR}/difexp_significant_females_1t_vs_2t.tsv")

    prater = load_prater()
    symbols = load_symbols()

    all_genes = sorted(sig_12 | sig_m | sig_f)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DEGs sex-stratified"

    headers = [
        "Entrez ID", "Symbol", "Gene Name", "Chr",
        "1_2 logFC", "1_2 FDR", "1_2 sig",
        "1_2_m logFC", "1_2_m FDR", "1_2_m sig",
        "1_2_f logFC", "1_2_f FDR", "1_2_f sig",
        "Venn group", "Near-miss", "Candidate sex-specific",
        "Prater logFC", "Prater padj",
    ]

    header_font = Font(bold=True, size=10, name="Calibri")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4",
                              fill_type="solid")
    header_font_white = Font(bold=True, size=10, name="Calibri",
                             color="FFFFFF")
    thin_border = Border(
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    italic_font = Font(italic=True, size=9, name="Calibri", color="666666")
    bold_font = Font(bold=True, size=9, name="Calibri", color="1A237E")
    normal_font = Font(size=9, name="Calibri")
    sig_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9",
                           fill_type="solid")
    truly_fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0",
                             fill_type="solid")
    near_miss_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5",
                                 fill_type="solid")

    for row_idx, gene in enumerate(all_genes, 2):
        in_12 = gene in sig_12
        in_m = gene in sig_m
        in_f = gene in sig_f

        r_12 = de_12.get(gene)
        r_m = de_m.get(gene)
        r_f = de_f.get(gene)

        sym_info = symbols.get(gene, {})
        symbol = sym_info.get("symbol", "")
        genename = sym_info.get("genename", "")
        chrom = sym_info.get("chr", "")

        logfc_12 = float(r_12["logFC"]) if r_12 else None
        fdr_12 = float(r_12["adj.P.Val"]) if r_12 else None
        logfc_m = float(r_m["logFC"]) if r_m else None
        fdr_m = float(r_m["adj.P.Val"]) if r_m else None
        logfc_f = float(r_f["logFC"]) if r_f else None
        fdr_f = float(r_f["adj.P.Val"]) if r_f else None

        # Venn group
        if in_m and in_f:
            venn = "shared"
        elif in_m and not in_f:
            venn = "male_only"
        elif in_f and not in_m:
            venn = "female_only"
        else:
            venn = "1_2_only"

        # Near-miss analysis
        near_miss_label = ""
        is_nm = False
        if venn == "male_only" and r_f and is_near_miss(r_f):
            near_miss_label = "near-miss in F"
            is_nm = True
        elif venn == "female_only" and r_m and is_near_miss(r_m):
            near_miss_label = "near-miss in M"
            is_nm = True

        # Candidate sex-specific
        truly_ss = False
        if venn == "male_only" and not is_nm:
            truly_ss = True
        elif venn == "female_only" and not is_nm:
            truly_ss = True

        # Prater
        p = prater.get(gene)
        p_logfc = p["log2FC"] if p else None
        p_padj = p["padj"] if p else None

        # Write row
        values = [
            gene, symbol, genename, chrom,
            logfc_12, fdr_12, "yes" if in_12 else "",
            logfc_m, fdr_m, "yes" if in_m else "",
            logfc_f, fdr_f, "yes" if in_f else "",
            venn, near_miss_label, "yes" if truly_ss else "",
            p_logfc, p_padj,
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border

            if truly_ss:
                cell.font = bold_font
                cell.fill = truly_fill
            elif is_nm:
                cell.font = italic_font
                cell.fill = near_miss_fill
            else:
                cell.font = normal_font

            # Format numbers
            if isinstance(val, float):
                if col_idx in (6, 9, 12, 18):  # FDR columns
                    cell.number_format = '0.00E+00'
                else:  # logFC columns
                    cell.number_format = '0.00'

    # Column widths
    col_widths = {
        1: 10, 2: 12, 3: 35, 4: 6,
        5: 10, 6: 10, 7: 7,
        8: 10, 9: 10, 10: 7,
        11: 10, 12: 10, 13: 7,
        14: 13, 15: 16, 16: 16,
        17: 12, 18: 12,
    }
    for col, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(all_genes) + 1}"
    ws.freeze_panes = "E2"

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    summary = [
        ["Category", "Count"],
        ["Total genes in table", len(all_genes)],
        ["Significant in 1_2 (combined)", len(sig_12)],
        ["Significant in 1_2_m (males)", len(sig_m)],
        ["Significant in 1_2_f (females)", len(sig_f)],
        [""],
        ["Venn groups:"],
        ["  Shared (in both M and F)", sum(1 for g in all_genes if g in sig_m and g in sig_f)],
        ["  Male-only", sum(1 for g in all_genes if g in sig_m and g not in sig_f)],
        ["  Female-only", sum(1 for g in all_genes if g in sig_f and g not in sig_m)],
        ["  1_2 only (not in M or F)", sum(1 for g in all_genes if g in sig_12 and g not in sig_m and g not in sig_f)],
        [""],
        ["Near-miss (italic):", sum(1 for g in all_genes
            if (g in sig_m and g not in sig_f and g in de_f and is_near_miss(de_f[g])) or
               (g in sig_f and g not in sig_m and g in de_m and is_near_miss(de_m[g])))],
        ["Candidate sex-specific (bold):", sum(1 for g in all_genes
            if (g in sig_m and g not in sig_f and (g not in de_f or not is_near_miss(de_f[g]))) or
               (g in sig_f and g not in sig_m and (g not in de_m or not is_near_miss(de_m[g]))))],
        [""],
        ["In Prater DEGs:", sum(1 for g in all_genes if g in prater)],
        [""],
        ["Thresholds:"],
        ["  FDR", FDR_THRESH],
        ["  |logFC|", LOGFC_THRESH],
        ["  Near-miss FDR", NEAR_MISS_FDR],
        ["  Near-miss |logFC|", NEAR_MISS_LOGFC],
    ]
    for row_idx, row in enumerate(summary, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = Font(bold=True)
    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 12

    wb.save(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"Total genes: {len(all_genes)}")


if __name__ == "__main__":
    main()
